"""Build the PPC Master Sheet template (one tab per project).

Run:  python ppc-master-sheet/build_master_sheet.py
Output: ppc-master-sheet/PPC_Master_Sheet.xlsx

Layout of every project tab
  Row 1  Project banner
  Row 2  Section groups (colour-coded)
  Row 3  Column headers  (never rename; the AI agent reads these)
  Row 4  What to enter / example (guide row, grey italic)
  Row 5  MASTER PROFILE: the current, always-up-to-date state of the project
  Row 6+ UPDATE LOG: one new row per update. Fill Date + Updated By + Entry Type,
         then write the update under the proper column. Never delete old rows.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).with_name("PPC_Master_Sheet.xlsx")
FONT = "Arial"
LOG_ROWS = 300  # pre-formatted update-log rows per project tab

# (section, colour, [(header, width, guide/example, dropdown-key or None)])
SECTIONS = [
    ("ENTRY INFO", "404040", [
        ("Entry Date", 13, "DD-MMM-YYYY. Profile row = last profile refresh date", None),
        ("Updated By (Team / AI Agent)", 16, "e.g. Rahul / AI Agent", None),
        ("Entry Type", 18, "Pick from list", "entry_type"),
    ]),
    ("PROJECT DETAILS", "1F4E78", [
        ("Name of the Project / Website", 26, "e.g. Summit Cleaning - summitcleaning.com", None),
        ("Platform (Google / Meta)", 14, "Google / Meta / Google + Meta", "platform"),
        ("Responsible Employee", 18, "Account owner name", None),
        ("Client Name", 18, "Client contact person + company", None),
        ("Account ID", 18, "Google Ads CID (xxx-xxx-xxxx) / Meta Ad Account ID", None),
        ("Project Stage", 20, "Pick from list", "stage"),
    ]),
    ("ACCESS & TRACKING TOOLS", "2E75B6", [
        ("Google Analytics Account", 20, "GA4 property ID + access status", None),
        ("Google Tag Manager", 18, "GTM-XXXXXXX + access status", None),
        ("Merchant Center Access", 16, "MC ID / N.A.", None),
        ("GBP Access", 16, "GBP name / access status", None),
        ("GSC Access", 16, "Property URL / access status", None),
        ("Microsoft Clarity", 16, "Project ID / access status", None),
        ("CRM Credential", 20, "CRM name + VAULT reference (never plain passwords)", None),
        ("CMS Access", 20, "WordPress/Shopify etc. + VAULT reference", None),
        ("Meta Business Manager / Pixel", 20, "BM ID, Pixel ID, access status (Meta projects)", None),
    ]),
    ("COMMERCIAL & BUDGET", "548235", [
        ("Client Payment Mode", 14, "Client pays Google/Meta by card / invoice / via agency", None),
        ("Package", 16, "Agency package name + fee", None),
        ("Start Date", 13, "Project start date", None),
        ("Monthly Date", 13, "Monthly cycle / billing date", None),
        ("Confirmation Received Date", 15, "Date client confirmed the project", None),
        ("Monthly Budget", 14, "Ad spend / month (with currency)", None),
        ("Daily Budget", 13, "Ad spend / day (with currency)", None),
        ("Current Bidding Strategy", 20, "e.g. Maximise Conversions, tCPA 40", None),
    ]),
    ("REPORTING CALENDAR", "BF8F00", [
        ("Weekly Reporting Date", 14, "Day of week, e.g. Monday", "weekday"),
        ("Weekly Conversion Reporting Date", 16, "Every Friday", "weekday"),
        ("Monthly Reporting Date", 14, "e.g. 3rd of every month", None),
        ("Client Last Meeting Date", 14, "Date of last call/meeting", None),
        ("Meeting Frequency", 14, "Bi-weekly / Monthly", "meeting"),
        ("Next Meeting Date", 14, "Scheduled date", None),
    ]),
    ("BUSINESS OVERVIEW & QUESTIONNAIRE", "7030A0", [
        ("Business Overview / Service Related Questionnaire", 40, "Summary: what they sell, where, to whom, avg order/job value", None),
        ("Full Service Questionnaire", 30, "Link to filled questionnaire + date received", None),
        ("Services / Products to Promote (Priority)", 30, "Service clusters in priority order", None),
        ("Target Locations", 22, "Cities / radius / countries", None),
    ]),
    ("AUDIENCE", "C55A11", [
        ("Primary Audiences", 30, "Main buyer persona", None),
        ("Secondary Audiences", 30, "Other buyers", None),
        ("Excluded Audiences", 26, "Who we must NOT target", None),
        ("Negative Terms That Exclude", 30, "Negative keywords / themes (e.g. jobs, free, DIY)", None),
        ("Jobs To Be Done for Audience", 30, "What the customer is trying to achieve", None),
        ("Use Cases / Scenarios (Audiences Looking for Services)", 34, "Situations that trigger the search", None),
        ("Problems & Pain Points", 30, "Pains the ads/landing page must address", None),
    ]),
    ("COMPETITION & POSITIONING", "833C0B", [
        ("Competitive Landscape", 32, "Top competitors + what they run (ads, offers, LPs)", None),
        ("Differentiation", 28, "Why choose the client over competitors", None),
        ("Proof Points or USPs", 28, "Reviews, years, certifications, guarantees (must be verifiable)", None),
        ("Auction Insights / Competitor Ads Notes", 32, "Impression share vs competitors, new entrants", None),
    ]),
    ("CONVERSION & TRACKING", "C00000", [
        ("Conversion Goal", 26, "Calls / forms / WhatsApp / purchases + target CPL/ROAS", None),
        ("Current Conversion Tracking Set Up Details", 36, "What is tracked, how (GTM/GA4/CMS/API), primary vs secondary", None),
        ("Tracking Status", 14, "Pick from list", "tracking"),
    ]),
    ("CAMPAIGN WORK & OPTIMIZATION", "00B0F0", [
        ("Keyword Research / Clusters Status", 26, "Sent / Approved on DATE + notes", None),
        ("Ad Assets Status (Copy, Images, Video)", 26, "Sent / Approved on DATE + notes", None),
        ("Campaign Structure / Live Campaigns", 30, "Campaign names + types (Search, PMax, Demand Gen, Meta)", None),
        ("Daily Work Update", 36, "Brief daily work: negatives added, new KWs from search terms, bids...", None),
        ("Negative Keywords Added", 26, "Terms + match type + level", None),
        ("New Keywords Added (from Search Terms)", 26, "Terms + match type + ad group", None),
        ("Landing Page / CRO Notes", 30, "LP vs competitor LP, CRO changes, new LP requests, QS impact", None),
        ("Last Optimization (Last 3 Months)", 40, "Rolling summary of changes in last 90 days (profile row)", None),
    ]),
    ("PERFORMANCE SNAPSHOT", "375623", [
        ("Period", 16, "e.g. 15-21 Sep 2026", None),
        ("Spend", 11, "Number", None),
        ("Impressions", 11, "Number", None),
        ("Clicks", 10, "Number", None),
        ("CTR", 9, "Auto = Clicks / Impressions", None),
        ("Conversions", 11, "Number (platform-reported)", None),
        ("Qualified Leads (Client Confirmed)", 13, "Number from CRM / client", None),
        ("CPL / CPA", 11, "Auto = Spend / Conversions", None),
        ("Conv. Rate", 10, "Auto = Conversions / Clicks", None),
        ("Calls / Forms / WhatsApp Split", 20, "e.g. 12 calls / 8 forms / 5 WA", None),
    ]),
    ("CLIENT COMMUNICATION", "FF6699", [
        ("Communication Channel", 14, "Pick from list", "channel"),
        ("Last Client Update or Comments or Feedback (Last 3 Months)", 44, "EVERY client message / feedback / approval, with date", None),
        ("Our Reply / Action Taken", 36, "What we replied and did", None),
        ("Client Complaint / Issue", 28, "Complaint details (if any)", None),
        ("Complaint Status", 13, "Pick from list", "complaint"),
        ("Meeting Notes (MoM)", 36, "Minutes of meeting + agreed actions", None),
    ]),
    ("RISK", "FF0000", [
        ("Risk Status", 12, "Green / Amber / Red", "risk"),
        ("Risk Reason", 30, "Why at risk (performance, budget, unhappy client, payment...)", None),
    ]),
    ("REPORTS & DRIVE LINKS", "3A3838", [
        ("Weekly Report Folder (Same Google Sheet Updated Weekly)", 30, "Drive link", None),
        ("Monthly Reporting Folder", 30, "Drive link", None),
        ("Media Doc", 30, "Drive link (images, videos, logos)", None),
        ("Conversion Reporting Format", 30, "Sheet link", None),
        ("Working Doc", 30, "Sheet link", None),
        ("Keyword File", 30, "Sheet link", None),
        ("Report / File Sent Link", 30, "Link of the report sent in this entry", None),
    ]),
    ("NEXT STEPS", "002060", [
        ("Next Action Plan or Recommendation", 40, "What we do next and why", None),
        ("Action Owner", 14, "Who", None),
        ("Due Date", 12, "When", None),
        ("Action Status", 13, "Pick from list", "action"),
    ]),
]

DROPDOWNS = {
    "entry_type": ["Master Profile", "Onboarding", "Access Update", "Keyword Research",
                   "Ad Assets", "Campaign Build", "Tracking Setup", "Daily Update",
                   "Optimization", "Weekly Report", "Weekly Conversion Report",
                   "Monthly Report", "Client Communication", "Client Complaint",
                   "Meeting", "Competitor / Auction Review", "Landing Page / CRO",
                   "Budget / Billing", "Risk Update"],
    "platform": ["Google", "Meta", "Google + Meta"],
    "stage": ["1-Questionnaire Sent", "2-Access Pending", "3-Keyword Research",
              "4-Ad Assets", "5-Campaign Build & Tracking", "6-Live & Optimizing",
              "Paused", "Churned"],
    "weekday": ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"],
    "meeting": ["Weekly", "Bi-weekly", "Monthly", "On request"],
    "tracking": ["Verified", "Partial", "Broken", "Not Set Up"],
    "channel": ["Basecamp", "WhatsApp", "Telegram", "Email", "Meeting / Call", "Other"],
    "complaint": ["Open", "In Progress", "Resolved"],
    "risk": ["Green", "Amber", "Red"],
    "action": ["To Do", "In Progress", "Waiting on Client", "Done"],
}

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def columns():
    """Flat list of (section, colour, header, width, guide, dropdown)."""
    return [(s, c, *f) for s, c, fields in SECTIONS for f in fields]


def col_of(header):
    for i, col in enumerate(columns(), start=1):
        if col[2] == header:
            return get_column_letter(i)
    raise KeyError(header)


def light(hex_colour):
    """Pale tint of a section colour for data-row banding."""
    r, g, b = (int(hex_colour[i:i + 2], 16) for i in (0, 2, 4))
    return "".join(f"{int(v + (255 - v) * 0.85):02X}" for v in (r, g, b))


def build_lists_sheet(wb):
    ws = wb.create_sheet("Lists")
    for i, (key, values) in enumerate(DROPDOWNS.items(), start=1):
        letter = get_column_letter(i)
        ws[f"{letter}1"] = key
        ws[f"{letter}1"].font = Font(name=FONT, bold=True)
        for r, v in enumerate(values, start=2):
            ws[f"{letter}{r}"] = v
            ws[f"{letter}{r}"].font = Font(name=FONT)
        ws.column_dimensions[letter].width = 26
    ws.sheet_state = "hidden"


def list_ref(key):
    idx = list(DROPDOWNS).index(key) + 1
    letter = get_column_letter(idx)
    return f"=Lists!${letter}$2:${letter}${len(DROPDOWNS[key]) + 1}"


def build_project_tab(wb, title, profile=None):
    ws = wb.create_sheet(title)
    cols = columns()
    last = get_column_letter(len(cols))
    last_row = 5 + LOG_ROWS

    # Row 1 banner
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = (f"{title.upper()}  |  Row 5 = MASTER PROFILE (keep current)  |  "
                "Row 6+ = UPDATE LOG (add a new row for every update, newest at the bottom, never delete)")
    ws["A1"].font = Font(name=FONT, bold=True, size=13, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="0B2545")
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26

    # Row 2 section bands, Row 3 headers, Row 4 guide
    col = 1
    for section, colour, fields in SECTIONS:
        start, end = col, col + len(fields) - 1
        if end > start:
            ws.merge_cells(start_row=2, start_column=start, end_row=2, end_column=end)
        cell = ws.cell(row=2, column=start, value=section)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=colour)
        cell.alignment = CENTER
        for c in range(start, end + 1):
            ws.cell(row=2, column=c).fill = PatternFill("solid", fgColor=colour)
        col = end + 1

    for i, (_, colour, header, width, guide, _dd) in enumerate(cols, start=1):
        letter = get_column_letter(i)
        h = ws.cell(row=3, column=i, value=header)
        h.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        h.fill = PatternFill("solid", fgColor=colour)
        h.alignment = CENTER
        h.border = BORDER
        g = ws.cell(row=4, column=i, value=guide)
        g.font = Font(name=FONT, italic=True, color="595959", size=9)
        g.fill = PatternFill("solid", fgColor="F2F2F2")
        g.alignment = WRAP_TOP
        g.border = BORDER
        ws.column_dimensions[letter].width = width
        tint = PatternFill("solid", fgColor=light(colour))
        for r in range(5, last_row + 1):
            c = ws.cell(row=r, column=i)
            c.font = Font(name=FONT, size=10, bold=(r == 5))
            c.alignment = WRAP_TOP
            c.border = BORDER
            if r == 5:
                c.fill = PatternFill("solid", fgColor="FFF2CC")
            elif r % 2 == 0:
                c.fill = tint
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 58
    ws.row_dimensions[4].height = 48
    ws.row_dimensions[5].height = 60

    ws.cell(row=5, column=3).comment = Comment(
        "MASTER PROFILE row. Overwrite these cells when something changes "
        "(budget, bidding, access, risk...). Record every change as a log row too.", "PPC Agent")

    # Auto metrics (profile + log rows)
    sp, im, cl, cv = (col_of(h) for h in ("Spend", "Impressions", "Clicks", "Conversions"))
    for r in range(5, last_row + 1):
        ws[f"{col_of('CTR')}{r}"] = f'=IF(AND(ISNUMBER({cl}{r}),N({im}{r})>0),{cl}{r}/{im}{r},"")'
        ws[f"{col_of('CPL / CPA')}{r}"] = f'=IF(AND(ISNUMBER({sp}{r}),N({cv}{r})>0),{sp}{r}/{cv}{r},"")'
        ws[f"{col_of('Conv. Rate')}{r}"] = f'=IF(AND(ISNUMBER({cv}{r}),N({cl}{r})>0),{cv}{r}/{cl}{r},"")'
        ws[f"{col_of('CTR')}{r}"].number_format = "0.00%"
        ws[f"{col_of('Conv. Rate')}{r}"].number_format = "0.00%"
        ws[f"{col_of('CPL / CPA')}{r}"].number_format = "#,##0.00"
        ws[f"{sp}{r}"].number_format = "#,##0.00"
        ws[f"A{r}"].number_format = "DD-MMM-YYYY"

    # Dropdowns
    for i, (*_, dd) in enumerate(cols, start=1):
        if dd:
            dv = DataValidation(type="list", formula1=list_ref(dd), allow_blank=True)
            letter = get_column_letter(i)
            dv.add(f"{letter}5:{letter}{last_row}")
            ws.add_data_validation(dv)

    # Profile row
    ws["C5"] = "Master Profile"
    for header, value in (profile or {}).items():
        ws[f"{col_of(header)}5"] = value

    ws.freeze_panes = "E5"  # keep Entry info + Project name + headers visible
    ws.auto_filter.ref = f"A3:{last}{last_row}"
    ws.sheet_view.zoomScale = 90
    return ws


def build_index(wb, projects):
    ws = wb.create_sheet("Project Index", 0)
    headers = [
        ("Project Tab Name (exact)", 26), ("Name of the Project / Website", 30),
        ("Platform (Google / Meta)", 14), ("Responsible Employee", 18), ("Client Name", 20),
        ("Account ID", 16), ("Project Stage", 20), ("Monthly Budget", 14),
        ("Weekly Reporting Date", 14), ("Weekly Conversion Reporting Date", 16),
        ("Monthly Reporting Date", 14), ("Risk Status", 11),
        ("Last Log Entry Date", 14), ("Next Action Plan or Recommendation", 40),
    ]
    ws.merge_cells("A1:N1")
    ws["A1"] = ("PPC MASTER SHEET - PROJECT INDEX  |  Type the exact tab name in column A; "
                "everything else fills automatically from that project tab's Row 5")
    ws["A1"].font = Font(name=FONT, bold=True, size=13, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="0B2545")
    ws.row_dimensions[1].height = 26
    for i, (h, w) in enumerate(headers, start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = CENTER
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[3].height = 44
    ws["A2"] = "Blue text = type here. Black = auto formulas (do not edit)."
    ws["A2"].font = Font(name=FONT, italic=True, color="595959", size=9)

    for r in range(4, 4 + 60):  # room for 60 projects
        a = ws.cell(row=r, column=1)
        a.font = Font(name=FONT, color="0000FF")
        a.border = BORDER
        for i, (h, _) in enumerate(headers[1:], start=2):
            c = ws.cell(row=r, column=i)
            c.font = Font(name=FONT, size=10)
            c.border = BORDER
            c.alignment = WRAP_TOP
            if h == "Last Log Entry Date":
                rng = f'INDIRECT("\'"&$A{r}&"\'!A6:A{5 + LOG_ROWS}")'
                c.value = f'=IF($A{r}="","",IFERROR(IF(MAX({rng})=0,"",MAX({rng})),""))'
                c.number_format = "DD-MMM-YYYY"
            else:
                c.value = (f'=IF($A{r}="","",IFERROR(INDIRECT("\'"&$A{r}&"\'!{col_of(h)}5")&"",""))')
    for r, name in enumerate(projects, start=4):
        ws.cell(row=r, column=1, value=name)
    ws.freeze_panes = "B4"
    return ws


def build_guide(wb):
    ws = wb.create_sheet("How To Use", 1)
    ws.column_dimensions["A"].width = 120
    lines = [
        ("PPC MASTER SHEET - HOW TO USE (Team + AI Agent)", True),
        ("", False),
        ("1. One tab per project. Copy the '_TEMPLATE' tab, rename it to the project name (e.g. 'Summit Cleaning') "
         "and add the name in 'Project Index' column A.", False),
        ("2. Row 3 = column headers. Never rename, delete or reorder them; the AI agent finds data by header name.", False),
        ("3. Row 4 = what to enter + example format.", False),
        ("4. Row 5 = MASTER PROFILE (yellow). Always holds the CURRENT truth: access, budget, bidding, audiences, "
         "USPs, tracking, links, risk, next action. Overwrite when something changes.", False),
        ("5. Row 6 onwards = UPDATE LOG. For every piece of work or client interaction add ONE new row at the bottom: "
         "Entry Date + Updated By + Entry Type, then write the details under the proper column(s). Never delete log rows.", False),
        ("6. Client communication: log EVERY message, feedback, approval or complaint (Basecamp, WhatsApp, Telegram, "
         "Email, Meeting) under 'Last Client Update or Comments or Feedback (Last 3 Months)' with the channel.", False),
        ("7. 'Last Optimization (Last 3 Months)' and 'Last Client Update...' in Row 5 are rolling 90-day summaries. "
         "Refresh them at every monthly report.", False),
        ("8. Performance Snapshot: enter Spend / Impressions / Clicks / Conversions for weekly & monthly report rows. "
         "CTR, CPL and Conv. Rate calculate automatically.", False),
        ("9. Security: do NOT paste passwords. Store credentials in the agency password manager and write only the "
         "vault item name in CRM Credential / CMS Access.", False),
        ("10. Risk Status: Green = on track; Amber = KPI off target 2+ weeks or client concern; Red = complaint, "
         "tracking broken, spend issue or churn risk. Red must have an Action Owner + Due Date.", False),
        ("", False),
        ("Colour code: dark banner = section; yellow row = master profile; grey italic = guide; blue text in index = input.", False),
    ]
    for r, (text, bold) in enumerate(lines, start=1):
        c = ws.cell(row=r, column=1, value=text)
        c.font = Font(name=FONT, bold=bold, size=14 if bold else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")


def main():
    wb = Workbook()
    wb.remove(wb.active)
    build_project_tab(wb, "_TEMPLATE")
    build_project_tab(wb, "Summit Cleaning", profile={
        "Name of the Project / Website": "Summit Cleaning",
        "Monthly Reporting Folder": "https://drive.google.com/drive/folders/1-IGAZ_1TWu6hhL5nTUQ2S3lRA46bk1gd",
        "Media Doc": "https://drive.google.com/drive/folders/1sdqNMqjWd_C4xaz8oJ0SjQqffzpeCYRE?usp=drive_link",
        "Conversion Reporting Format": "https://docs.google.com/spreadsheets/d/1aso6ehzhsEiwWslTu0zgwIodOsFwl6a3o_-cKLkHBoE/edit?gid=1837517691#gid=1837517691",
        "Working Doc": "https://docs.google.com/spreadsheets/d/1L-3Vsn272x2WBcjAqTtffOsufC7_MxWuuFDq48ZEjfQ/edit?gid=1727187873#gid=1727187873",
        "Keyword File": "https://docs.google.com/spreadsheets/d/173R_v8cuwZ2JxPqV-U5dQrxkhMGp1Y6F3hjRTefeBNw/edit?gid=1257750804#gid=1257750804",
    })
    build_index(wb, ["Summit Cleaning"])
    build_guide(wb)
    build_lists_sheet(wb)
    wb.active = 0
    wb.save(OUT)
    print(f"Saved {OUT} ({len(columns())} columns per project tab)")


if __name__ == "__main__":
    main()
